"""数据驱动：从 CSV / Excel 加载参数化数据

- CSV 使用标准库，零依赖；Excel 需要 openpyxl（可选依赖 atf[excel]）
- 每行数据转换为一组 flat 变量，用例按行重复执行
- 数据值统一按字符串处理，避免数字列在请求体中意外变成数字类型；
  断言时 eq/ne 会做字符串回退比较（见 assertion.py）
"""
import csv
from pathlib import Path

from .errors import ATFError


def resolve_data_path(case: dict, datas: str) -> Path:
    """把 datas 解析为绝对路径：绝对路径原样返回，相对路径基于用例文件所在目录"""
    path = Path(datas)
    if path.is_absolute():
        return path
    case_file = case.get("__file__")
    base = Path(case_file).parent if case_file else Path.cwd()
    return base / path


def load_data_file(path) -> list:
    """加载数据文件，返回逐行数据（列表字典）"""
    path = Path(path)
    if not path.exists():
        raise ATFError(f"数据文件不存在：{path}")
    suffix = path.suffix.lower()
    if suffix == ".csv":
        rows = _load_csv(path)
    elif suffix in (".xlsx", ".xlsm"):
        rows = _load_excel(path)
    elif suffix == ".xls":
        raise ATFError("不支持旧版 .xls 格式，请另存为 .xlsx 或 .csv 后再试")
    else:
        raise ATFError(f"不支持的数据文件格式：{suffix}（仅支持 .csv / .xlsx）")
    if not rows:
        raise ATFError(f"数据文件没有可用数据行：{path}")
    return rows


def _load_csv(path) -> list:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        return [_convert_row(row) for row in reader]


def _load_excel(path) -> list:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ATFError("读取 Excel 需要 openpyxl，请先安装：pip install atf[excel]") from None
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        try:
            headers = [str(cell) if cell is not None else "" for cell in next(rows)]
        except StopIteration:
            return []
        records = []
        for values in rows:
            if values is None or all(value is None for value in values):
                continue
            records.append(
                {header: _convert(value) for header, value in zip(headers, values) if header}
            )
        return records
    finally:
        workbook.close()


def _convert_row(row: dict) -> dict:
    return {key: _convert(value) for key, value in row.items() if key}


def _convert(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value)
